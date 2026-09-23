# -*- coding: utf-8 -*-
"""把 B 成员填好的打分表转成 gold.json（封存）

为什么要有这一步：
    人工分数留在 Excel 里没法被程序读取，也没法「封存」。
    转成的 gold.json 是 D6 benchmark 的唯一标准答案，生成后即视为冻结：
    任何修改都必须重新打分并重新生成，不允许手改 json。

用法：
    python tools/import_gold.py

校验规则：
    - 一份报告要么 5 项全填，要么整份不填（整份不填的会明确列为「未纳入」，
      例如 S06 因超长系统评不了、不参与 benchmark，可以先不打）
    - 半份报告（填了 1~4 项）直接报错拒绝生成——半份 gold 会让评测失真
    - 每项得分必须在 0 ~ 该项满分之间
    - 判定必须是 hit / partial / miss
"""
import os
import sys
import json
import datetime

try:
    from openpyxl import load_workbook
except ImportError:
    print("缺少 openpyxl，请先安装： pip install openpyxl")
    sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "gold", "gold_人工打分表_空白.xlsx")
OUT = os.path.join(ROOT, "data", "gold", "gold.json")

MAXSCORE = {"r1": 15, "r2": 20, "r3": 25, "r4": 20, "r5": 20}
VERDICTS = {"hit", "partial", "miss"}


def cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    if isinstance(v, str):
        v = v.strip()
    return v


def read_seal(wb):
    """读封存记录。

    版式提示（踩过坑）：第 1 行是合并标题、第 2 行留空，字段从**第 3 行**才开始。
    表格下方可能还有「补填说明」等补充行，一并带出来——
    这样 gold.json 自己能说清来历（比如哪些字段是事后补填的）。
    """
    if "⑤ 封存记录" not in wb.sheetnames:
        return {}
    ws = wb["⑤ 封存记录"]
    info = {}
    for r in range(3, ws.max_row + 1):
        k = cell(ws, r, 1)
        v = cell(ws, r, 2)
        if k:
            info[str(k)] = v
    return info


def main():
    if not os.path.exists(XLSX):
        print(f"找不到打分表：{XLSX}")
        sys.exit(1)

    wb = load_workbook(XLSX, data_only=True)   # data_only：读公式的计算结果
    ws = wb["③ 逐项打分"]

    by_rep = {}        # rep -> 有效行
    blank_rep = {}     # rep -> 空行数（整份未填的候选）
    problems = []

    for r in range(2, ws.max_row + 1):
        rep = cell(ws, r, 2)
        iid = cell(ws, r, 3)
        if not rep or not iid:
            continue
        mx = MAXSCORE.get(iid)
        score = cell(ws, r, 6)
        verdict = cell(ws, r, 7)
        evidence = cell(ws, r, 8)
        note = cell(ws, r, 9)

        if score is None or score == "":
            # 空行先记下来：只有「整份都空」才算合法排除，零星空行仍是问题
            blank_rep[rep] = blank_rep.get(rep, 0) + 1
            by_rep.setdefault(rep, {})
            continue
        try:
            score = float(score)
        except (TypeError, ValueError):
            problems.append(f"第{r}行 {rep}/{iid}：得分不是数字（{score!r}）")
            continue
        if not (0 <= score <= mx):
            problems.append(f"第{r}行 {rep}/{iid}：得分 {score} 超出 0~{mx}")
            continue
        # verdict 必填（文档一直这么要求，旧实现却允许为空）
        if not verdict:
            problems.append(f"第{r}行 {rep}/{iid}：判定未填（必须 hit / partial / miss）")
            continue
        if verdict not in VERDICTS:
            problems.append(f"第{r}行 {rep}/{iid}：判定 {verdict!r} 不合法")
            continue

        by_rep.setdefault(rep, {})[iid] = {
            "score": score, "verdict": verdict,
            "evidence": evidence or "", "note": note or ""}

    # 整份未填 → 明确列为「未纳入」；填了一半 → 报错
    gold = {}
    excluded = []
    filled = 0
    for rep in sorted(by_rep):
        rec = by_rep[rep]
        if len(rec) == len(MAXSCORE):
            gold[rep] = {"items": rec,
                         "total": round(sum(v["score"] for v in rec.values()), 1)}
            filled += len(rec)
        elif len(rec) == 0:
            excluded.append(rep)          # 整份空着 = 主动不纳入，合法
        else:
            problems.append(f"{rep}：只填了 {len(rec)}/{len(MAXSCORE)} 项，"
                            f"半份报告不能进 gold（要么全填，要么整份空着）")

    print(f"完整填写 {filled} 项（{len(gold)} 份报告）")
    if excluded:
        print(f"未纳入 {len(excluded)} 份（整份未填）：{excluded}")
    if problems:
        print(f"\n发现 {len(problems)} 个问题，未生成 gold.json：")
        for p in problems[:20]:
            print("  -", p)
        if len(problems) > 20:
            print(f"  …… 还有 {len(problems)-20} 条")
        sys.exit(1)
    if not gold:
        print("没有任何一份报告完整填写，拒绝生成。")
        sys.exit(1)

    seal = read_seal(wb)
    # 评分者信息必须完整，否则这份 gold set 的公信力无从证明
    required = {
        "打分人姓名": seal.get("打分人姓名"),
        "打分日期": seal.get("打分日期"),
        "是否独立完成（未与主程讨论）": seal.get("是否独立完成（未与主程讨论）"),
        "是否全程未参考 AI 评分": seal.get("是否全程未参考 AI 评分"),
    }
    for k, v in required.items():
        if not v:
            problems.append(f"封存记录缺少「{k}」——无法证明这份 gold 是独立盲评的结果")
    if required["是否独立完成（未与主程讨论）"] not in (None, "", "是"):
        problems.append("「是否独立完成」不是「是」，该 gold set 不能用于对外指标")
    if required["是否全程未参考 AI 评分"] not in (None, "", "是"):
        problems.append("「是否全程未参考 AI 评分」不是「是」，该 gold set 不能用于对外指标")
    if problems:
        print(f"\n发现 {len(problems)} 个问题，未生成 gold.json：")
        for p in problems[:20]:
            print("  -", p)
        sys.exit(1)

    data = {
        "_note": "人工 gold set，非主程成员独立打分，生成后封存，禁止手改",
        "created_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "scorer": seal.get("打分人姓名") or "",
        "scorer_role": seal.get("专业 / 分工") or "",
        "scored_date": seal.get("打分日期") or "",
        "minutes_spent": seal.get("总耗时（分钟）") or "",
        "independent": seal.get("是否独立完成（未与主程讨论）") or "",
        "no_ai_reference": seal.get("是否全程未参考 AI 评分") or "",
        # 表下方的补充说明（例如"某几个字段是事后补填的"）原样带出，
        # 让 gold.json 自己交代来历，而不是靠口头解释
        "seal_note": seal.get("补填说明") or "",
        "source_sheet": os.path.basename(XLSX),
        "rubric": MAXSCORE,
        "excluded_reports": excluded,
        "reports": gold,
    }
    json.dump(data, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    print(f"\n✅ 已生成 {OUT}")
    for k, v in sorted(gold.items()):
        print(f"   {k}: {v['total']}")
    vals = [v["total"] for v in gold.values()]
    print(f"\n最高 {max(vals)} 最低 {min(vals)} 极差 {round(max(vals)-min(vals),1)}")
    if seal.get("是否独立完成（未与主程讨论）") != "是":
        print("\n⚠ 封存记录里「是否独立完成」不是「是」——这份 gold set 的公信力会受质疑，请确认。")


if __name__ == "__main__":
    main()
