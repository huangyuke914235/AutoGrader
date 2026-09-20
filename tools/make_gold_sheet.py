# -*- coding: utf-8 -*-
"""生成 gold set 人工打分表（空白模板）

给 B 成员（非主程）独立打分用。设计原则见 00-经验教训与防自欺纪律.md：
  - 验收标准不得由执行者自写
  - gold set 必须在 prompt 调优【之前】完成并封存，D6 才解封
  - 打分人全程不得看 AI 分数、不得与主程讨论

用法：
    python tools/make_gold_sheet.py
输出：
    data/gold/gold_人工打分表_空白.xlsx
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "data", "gold")

# 与 tools/make_demo.py、docs/cases 公布的案例完全一致的固定 rubric
ITEMS = [
    ("r1", "实验目的明确", 15, "开头明确写出本次实验的目的与要掌握的能力",
     "实验目的 / 旨在 / 掌握 / 目的"),
    ("r2", "环境与步骤", 20, "写清实验环境配置与可复现的操作步骤",
     "环境 / 步骤 / 安装 / 配置 / 命令"),
    ("r3", "核心实现", 25, "给出核心代码、模型结构或关键实现说明",
     "代码 / 实现 / 模型 / 算法 / 结构"),
    ("r4", "结果与数据", 20, "给出运行结果、截图、表格或实验数据",
     "结果 / 输出 / 截图 / 数据 / 表"),
    ("r5", "分析与总结", 20, "对结果进行分析讨论，并有总结或心得",
     "分析 / 总结 / 心得 / 讨论 / 结论"),
]
REPORTS = [f"S{i:02d}" for i in range(1, 11)]

# ---------- 样式 ----------
THIN = Side(style="thin", color="D8D6CF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_TITLE = Font(name="微软雅黑", size=15, bold=True, color="1F1F1D")
F_H2 = Font(name="微软雅黑", size=11, bold=True, color="185FA5")
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color="1F1F1D")
F_BODY = Font(name="微软雅黑", size=10, color="1F1F1D")
F_GRAY = Font(name="微软雅黑", size=10, color="5F5E5A")
FILL_HEAD = PatternFill("solid", fgColor="E6F1FB")
FILL_PRE = PatternFill("solid", fgColor="F5F4F0")   # 已预填、无需改
FILL_IN = PatternFill("solid", fgColor="FFFDF5")    # 需要人工填
FILL_WARN = PatternFill("solid", fgColor="FAEEDA")
AL_W = Alignment(wrap_text=True, vertical="top")
AL_C = Alignment(horizontal="center", vertical="center")
AL_CT = Alignment(horizontal="center", vertical="top", wrap_text=True)


def put(ws, row, col, value, font=F_BODY, fill=None, align=AL_W, border=True):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = BD
    return c


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


def tall(ws, row, text, base=17):
    """按文字长度估算行高，避免换行被截断"""
    lines = 1 + sum(max(1, len(p) // 52) for p in str(text).split("\n"))
    ws.row_dimensions[row].height = max(base, 16 * lines)


def sheet_intro(wb):
    ws = wb.create_sheet("① 说明与纪律")
    set_widths(ws, [26, 96])
    put(ws, 1, 1, "AutoGrader · gold set 人工打分表（空白模板）", F_TITLE, align=AL_C)
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 30
    put(ws, 2, 1, "打分对象", F_H2, FILL_HEAD)
    put(ws, 2, 2, "10 份脱敏实验报告 S01–S10（data/samples/），按 5 个评分点逐项打分，满分 100。", F_BODY)
    tall(ws, 2, "x" * 110)

    blocks = [
        ("一、这份表为什么要存在",
         "我们用它来检验 AutoGrader 的评分到底准不准。\n"
         "准确率不能由写代码的人自己宣布——那样想多少分就有多少分。\n"
         "所以这份人工分数是「标准答案（gold set）」，由没有参与代码和 prompt 调优的人独立给出。"),
        ("二、三条纪律（不可违反）",
         "1. 独立打分：全程不要和主程讨论任何一份报告该给多少分，不要问「系统给了几分」。\n"
         "2. 先于调优：这份表必须在任何 prompt 修改【之前】完成并封存。\n"
         "   封存后到 D6（benchmark）才解封，中间不允许回头改分。\n"
         "3. 不许偷看：打分期间不要打开系统的评阅结果、不要看 docs/cases 里的案例。"),
        ("三、打分流程",
         "1. 打开 data/samples/S01.txt，通读一遍（不要跳读）。\n"
         "2. 到「③ 逐项打分」表，给 S01 的 r1–r5 各打一个分，并在「依据」里抄一句原文或写清位置。\n"
         "3. 十份报告依次做完，共 50 行。\n"
         "4. 到「④ 总分汇总」核对每份总分（公式自动加总，不用手算）。\n"
         "5. 到「⑤ 封存记录」填写署名与时间，确认三项「是」后封存。"),
        ("四、判定与给分规则",
         "hit（命中）    → 给满分或接近满分（≥85%）。原文清楚写了这点要求的全部关键内容。\n"
         "partial（部分）→ 给满分的 40%–70%。提到了但说得不完整、只有标题没有实质内容、或只有现象没有分析。\n"
         "miss（未命中） → 给 0–20%。完全没写，或写的不是这回事。\n"
         "说明：百分比是参考，最终以你的判断为准——但每个分数都必须能在「依据」里指认出原文。\n"
         "特别注意：报告里只有小标题（比如只写了「运行结果」四个字）而没有实际内容，算 partial，不算 hit。"),
        ("五、填写要求",
         "· 「人工得分」必填，可以是小数（如 12.5）。\n"
         "· 「判定」从下拉选 hit / partial / miss。\n"
         "· 「依据」尽量填：抄一句原文片段，或写「第 3 节」「开头第 2 段」这类位置。\n"
         "   填不出来的分数，说明这个分给得没把握——那就重新看报告，或降到 partial。\n"
         "· 十份报告建议分 2–3 次做完，中途休息，避免越打越松（顺序效应）。\n"
         "· 全部打完前【不要】和主程讨论任何分数。"),
        ("六、常见疑问",
         "Q：拿不准该给 12 还是 14？\n"
         "A：给 13，然后在备注写「介于两者之间」。不要为了凑整数反复纠结，一致性比精度重要。\n\n"
         "Q：报告写得很烂，五项都该 0 吗？\n"
         "A：只要开头写了实验目的，r1 就该有分。逐项独立判断，不要因为整体印象差就全给 0。\n\n"
         "Q：可以改分吗？\n"
         "A：封存前可以。一旦在「⑤ 封存记录」确认，就不能再改——这是这份表可信的前提。"),
    ]
    r = 4
    for title, body in blocks:
        put(ws, r, 1, title, F_H2, FILL_HEAD, align=AL_W)
        put(ws, r, 2, body, F_BODY, FILL_IN)
        tall(ws, r, body)
        r += 1
    ws.freeze_panes = "A2"
    return ws


def sheet_rubric(wb):
    ws = wb.create_sheet("② 评分点定义")
    set_widths(ws, [10, 18, 10, 46, 34])
    put(ws, 1, 1, "评分点ID", F_HEAD, FILL_HEAD, AL_C)
    put(ws, 1, 2, "评分点名称", F_HEAD, FILL_HEAD, AL_C)
    put(ws, 1, 3, "满分", F_HEAD, FILL_HEAD, AL_C)
    put(ws, 1, 4, "判定标准", F_HEAD, FILL_HEAD, AL_C)
    put(ws, 1, 5, "命中特征词（参考）", F_HEAD, FILL_HEAD, AL_C)
    r = 2
    for iid, name, mx, crit, sig in ITEMS:
        put(ws, r, 1, iid, F_BODY, FILL_PRE, AL_C)
        put(ws, r, 2, name, F_BODY, FILL_PRE, AL_C)
        put(ws, r, 3, mx, F_BODY, FILL_PRE, AL_C)
        put(ws, r, 4, crit, F_BODY, FILL_PRE)
        put(ws, r, 5, sig, F_GRAY, FILL_PRE)
        tall(ws, r, crit, 20)
        r += 1
    put(ws, r, 1, "合计", F_HEAD, PatternFill("solid", fgColor="E1F5EE"), AL_C)
    put(ws, r, 2, "", F_HEAD, PatternFill("solid", fgColor="E1F5EE"), AL_C)
    put(ws, r, 3, f"=SUM(C2:C{r-1})", F_HEAD, PatternFill("solid", fgColor="E1F5EE"), AL_C)
    put(ws, r, 4, "与 docs/cases 公布案例、tools/make_demo.py 使用的标准完全一致",
        F_GRAY, PatternFill("solid", fgColor="E1F5EE"))
    put(ws, r, 5, "", F_GRAY, PatternFill("solid", fgColor="E1F5EE"))
    ws.freeze_panes = "A2"
    return ws


def sheet_items(wb):
    ws = wb.create_sheet("③ 逐项打分")
    set_widths(ws, [6, 10, 10, 16, 8, 12, 12, 46, 26])
    heads = ["序号", "报告ID", "评分点ID", "评分点名称", "满分",
             "人工得分\n(必填)", "判定\n(下拉)", "依据（原文摘录或位置，尽量填）", "备注"]
    for i, h in enumerate(heads, start=1):
        put(ws, 1, i, h, F_HEAD, FILL_HEAD, AL_C)
    ws.row_dimensions[1].height = 34

    r = 2
    n = 0
    for rep in REPORTS:
        for iid, name, mx, _c, _s in ITEMS:
            n += 1
            put(ws, r, 1, n, F_GRAY, FILL_PRE, AL_CT)
            put(ws, r, 2, rep, F_BODY, FILL_PRE, AL_CT)
            put(ws, r, 3, iid, F_BODY, FILL_PRE, AL_CT)
            put(ws, r, 4, name, F_BODY, FILL_PRE, AL_CT)
            put(ws, r, 5, mx, F_BODY, FILL_PRE, AL_CT)
            put(ws, r, 6, None, F_BODY, FILL_IN, AL_CT)   # 人工得分
            put(ws, r, 7, None, F_BODY, FILL_IN, AL_CT)   # 判定
            put(ws, r, 8, None, F_BODY, FILL_IN, AL_W)    # 依据
            put(ws, r, 9, None, F_GRAY, FILL_IN, AL_W)    # 备注
            ws.row_dimensions[r].height = 22

            dv_score = DataValidation(type="decimal", operator="between",
                                      formula1="0", formula2=str(mx),
                                      allow_blank=True, showErrorMessage=True)
            dv_score.errorTitle = "超出范围"
            dv_score.error = f"该项满分 {mx} 分，请输入 0 到 {mx} 之间的数字。"
            dv_score.promptTitle = "人工得分"
            dv_score.prompt = f"0 – {mx}，可填小数"
            ws.add_data_validation(dv_score)
            dv_score.add(ws.cell(row=r, column=6))

            dv_v = DataValidation(type="list", formula1='"hit,partial,miss"',
                                  allow_blank=True, showErrorMessage=True)
            dv_v.errorTitle = "只能选三项"
            dv_v.error = "请从下拉里选 hit / partial / miss。"
            ws.add_data_validation(dv_v)
            dv_v.add(ws.cell(row=r, column=7))
            r += 1
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:I{r-1}"
    return ws, r - 1


def sheet_total(wb, last_row):
    ws = wb.create_sheet("④ 总分汇总")
    set_widths(ws, [10, 20, 18, 20, 18, 18, 12, 34, 12])
    heads = ["报告ID", "r1 实验目的(15)", "r2 环境与步骤(20)", "r3 核心实现(25)",
             "r4 结果与数据(20)", "r5 分析与总结(20)", "总分", "整体印象（可选）", "耗时(分钟)"]
    for i, h in enumerate(heads, start=1):
        put(ws, 1, i, h, F_HEAD, FILL_HEAD, AL_C)
    ws.row_dimensions[1].height = 30

    ids = ["r1", "r2", "r3", "r4", "r5"]
    r = 2
    for rep in REPORTS:
        put(ws, r, 1, rep, F_BODY, FILL_PRE, AL_CT)
        for j, iid in enumerate(ids, start=2):
            f = (f'=SUMIFS(\'③ 逐项打分\'!$F$2:$F${last_row},'
                 f'\'③ 逐项打分\'!$B$2:$B${last_row},$A{r},'
                 f'\'③ 逐项打分\'!$C$2:$C${last_row},"{iid}")')
            put(ws, r, j, f, F_BODY, FILL_PRE, AL_CT)
        put(ws, r, 7, f"=SUM(B{r}:F{r})", F_HEAD, PatternFill("solid", fgColor="E1F5EE"), AL_CT)
        put(ws, r, 8, None, F_GRAY, FILL_IN, AL_W)
        put(ws, r, 9, None, F_GRAY, FILL_IN, AL_CT)
        ws.row_dimensions[r].height = 22
        r += 1

    put(ws, r + 1, 1, "说明", F_H2, FILL_HEAD, AL_C)
    put(ws, r + 1, 2,
        "本表各分数由公式从「③ 逐项打分」自动汇总，不需要手算。"
        "若某格显示 0，说明该项还没打分（或真的给了 0 分）。", F_GRAY, FILL_WARN, AL_W)
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=9)
    ws.freeze_panes = "B2"
    return ws


def sheet_seal(wb):
    ws = wb.create_sheet("⑤ 封存记录")
    set_widths(ws, [30, 34, 62])
    put(ws, 1, 1, "封存记录（打分完成后填写）", F_TITLE, align=AL_C)
    ws.merge_cells("A1:C1")
    ws.row_dimensions[1].height = 28

    rows = [
        ("打分人姓名", "", "请填真实姓名——benchmark 报告里要写「由非主程成员 XXX 独立打分」"),
        ("专业 / 分工", "", "例如：计算机科学与技术 / 产品与材料"),
        ("打分日期", "", "完成当天日期"),
        ("总耗时（分钟）", "", "用于说明这份 gold set 的认真程度"),
        ("是否独立完成（未与主程讨论）", "是", "必须是「是」，否则这份表作废"),
        ("是否全程未参考 AI 评分", "是", "必须是「是」"),
        ("是否确认不再修改任何分值", "是", "确认后即封存"),
        ("封存后由谁保管", "主程（只读，不得修改）", "主程保管但不得改动内容"),
        ("解封条件", "D6 跑 benchmark 当天解封", "解封后只读对比，仍不得改分"),
    ]
    r = 3
    for k, v, hint in rows:
        put(ws, r, 1, k, F_H2, FILL_HEAD, AL_W)
        put(ws, r, 2, v, F_BODY, FILL_IN, AL_CT)
        put(ws, r, 3, hint, F_GRAY, FILL_PRE, AL_W)
        tall(ws, r, hint, 22)
        if v == "是":
            dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(ws.cell(row=r, column=2))
        r += 1

    put(ws, r + 1, 1, "填写进度", F_H2, FILL_HEAD, AL_W)
    put(ws, r + 1, 2, "=COUNTA('③ 逐项打分'!F2:F51)&\" / 50\"", F_HEAD,
        PatternFill("solid", fgColor="E1F5EE"), AL_CT)
    put(ws, r + 1, 3, "显示已填得分条目数，全部完成应为 50 / 50", F_GRAY, FILL_PRE, AL_W)
    return ws


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    sheet_intro(wb)
    sheet_rubric(wb)
    _, last_row = sheet_items(wb)
    sheet_total(wb, last_row)
    sheet_seal(wb)
    out = os.path.join(OUT_DIR, "gold_人工打分表_空白.xlsx")
    wb.save(out)
    print(f"已生成：{out}")
    print(f"报告 {len(REPORTS)} 份 × 评分点 {len(ITEMS)} 项 = {len(REPORTS)*len(ITEMS)} 行待填")


if __name__ == "__main__":
    main()
